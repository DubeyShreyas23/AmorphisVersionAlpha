"""
Reorder DRDO_AMORPHIS_intro_2 Expanded.pptx per user spec, inserting the
two road-timeline slides from Amorphis-Timeline.pptx at position 4.

Current slides (1-based):
 1  Intro (AMORPHIS Space Systems)
 2  The Problem
 3  Our Solution
 4  Our Products (Competitive Edge)
 5  Our Vision
 6  Why SMAs (SMA explanation 1)
 7  SMA versatility (SMA explanation 2)
 8  WHY AMORPHIS – Tech dev
 9  WHY AMORPHIS – Sys eng
10  WHY AMORPHIS – Qualification
11  Why SMA wins – head-to-head / specs comparison
12  Defense Applications 1
13  Defense Applications 2
14  Strategic Independence
15  Product DEMOs
16  Five moats – WHY US
17  Obsessed Engineers – About Us
18  Closing (next satellite)

Target order:
 1  slide1  (Intro)
 2  slide17 (About Us / Obsessed Engineers)
 3  slide15 (Product Demos)
 4  [timeline slide 1  – inserted from Amorphis-Timeline.pptx]
 5  [timeline slide 2  – inserted from Amorphis-Timeline.pptx]
 6  slide11 (Specs vs EBAD / head-to-head)
 7  slide5  (Our Vision)
 8  slide4  (Our Products)
 9  slide12 (Defense Applications 1)
10  slide13 (Defense Applications 2)
11  slide14 (Strategic Independence)
12  slide16 (Five Moats – Why Us intro)
13  slide8  (WHY AMORPHIS 1 – Tech dev)
14  slide9  (WHY AMORPHIS 2 – Sys eng)
15  slide10 (WHY AMORPHIS 3 – Qualification)
16  slide6  (SMA explanation 1)
17  slide7  (SMA explanation 2)
18  slide2  (The Problem)
19  slide3  (Our Solution)
20  slide18 (Closing)
"""

import copy, os
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from lxml import etree

DRDO_PATH     = r'C:\Startup Idea Amorphis\DRDO_AMORPHIS_intro_2 Expanded.pptx'
TIMELINE_PATH = r'C:\Startup Idea Amorphis\Amorphis-Timeline.pptx'
OUT_PATH      = r'C:\Startup Idea Amorphis\DRDO_AMORPHIS_reordered.pptx'

# ── STEP 1: Reorder existing slides ─────────────────────────────────────────
drdo     = Presentation(DRDO_PATH)
timeline = Presentation(TIMELINE_PATH)

print(f'DRDO:     {len(drdo.slides)} slides,  {drdo.slide_width.inches:.3f}" x {drdo.slide_height.inches:.3f}"')
print(f'Timeline: {len(timeline.slides)} slides,  {timeline.slide_width.inches:.3f}" x {timeline.slide_height.inches:.3f}"')

sld_id_lst = drdo.element.find(qn('p:sldIdLst'))
orig_sld_ids = list(sld_id_lst)   # keep copy of original 0-based list

# New order for EXISTING slides before timeline insertion (0-based indices):
#   pos  0 → original slide 0  (Intro)
#   pos  1 → original slide 16 (About Us)
#   pos  2 → original slide 14 (Product Demos)
#   pos  3 → original slide 10 (Specs)
#   pos  4 → original slide  4 (Vision)
#   pos  5 → original slide  3 (Products)
#   pos  6 → original slide 11 (Defense 1)
#   pos  7 → original slide 12 (Defense 2)
#   pos  8 → original slide 13 (Strategic Independence)
#   pos  9 → original slide 15 (Five Moats)
#   pos 10 → original slide  7 (WHY AMORPHIS 1)
#   pos 11 → original slide  8 (WHY AMORPHIS 2)
#   pos 12 → original slide  9 (WHY AMORPHIS 3)
#   pos 13 → original slide  5 (SMA 1)
#   pos 14 → original slide  6 (SMA 2)
#   pos 15 → original slide  1 (Problem)
#   pos 16 → original slide  2 (Solution)
#   pos 17 → original slide 17 (Closing)

EXISTING_ORDER = [0, 16, 14, 10, 4, 3, 11, 12, 13, 15, 7, 8, 9, 5, 6, 1, 2, 17]

reordered = [orig_sld_ids[i] for i in EXISTING_ORDER]
# Replace the sldIdLst contents with the new order
for sld in list(sld_id_lst):
    sld_id_lst.remove(sld)
for sld in reordered:
    sld_id_lst.append(sld)

print('Reordered existing slides.')

# ── STEP 2: Copy timeline slides in, scaling to DRDO dimensions ──────────────
# DRDO:     10.0"  x 5.625"  →  9144000 x 5143500 EMU
# Timeline: 13.333" x 7.5"   → 12192000 x 6858000 EMU
# Scale factor: 9144000 / 12192000 = 0.75

SCALE = drdo.slide_width.emu / timeline.slide_width.emu
print(f'Scale factor timeline → DRDO: {SCALE:.6f}')

COORD_ATTRS = ('cx', 'cy', 'x', 'y', 'l', 't', 'r', 'b')

def scale_xfrm(el, scale):
    """Recursively scale all <a:off> and <a:ext> attributes in an element."""
    tag = el.tag.split('}')[-1] if '}' in el.tag else el.tag
    if tag in ('off', 'ext', 'chOff', 'chExt'):
        for attr in ('cx', 'cy', 'x', 'y'):
            if el.get(attr) is not None:
                el.set(attr, str(int(int(el.get(attr)) * scale)))
    for child in el:
        scale_xfrm(child, scale)

def find_blank_layout(prs):
    """Return the most 'blank' slide layout in the presentation."""
    for layout in prs.slide_layouts:
        if 'blank' in layout.name.lower():
            return layout
    return prs.slide_layouts[-1]   # fallback

blank_layout = find_blank_layout(drdo)

def copy_slide_at(src_prs, src_idx, tgt_prs, insert_pos, scale=1.0):
    """
    Copy slide[src_idx] from src_prs to tgt_prs at insert_pos (0-based).
    Scales all shape coordinates by `scale`.
    """
    src_slide = src_prs.slides[src_idx]

    # Add a blank slide (appended at the end first)
    new_slide = tgt_prs.slides.add_slide(blank_layout)

    # Replace the spTree with source content
    sp_tree  = new_slide.shapes._spTree
    src_tree = src_slide.shapes._spTree
    for child in list(sp_tree):
        sp_tree.remove(child)
    for child in src_tree:
        new_child = copy.deepcopy(child)
        if scale != 1.0:
            scale_xfrm(new_child, scale)
        sp_tree.append(new_child)

    # Also copy slide background (bg element)
    src_bg_elem = src_slide._element.find(qn('p:bg'))
    if src_bg_elem is not None:
        tgt_bg = new_slide._element.find(qn('p:bg'))
        if tgt_bg is not None:
            new_slide._element.remove(tgt_bg)
        new_slide._element.insert(2, copy.deepcopy(src_bg_elem))

    # The new slide was appended to the end — move it to insert_pos
    sld_id_lst_now = tgt_prs.element.find(qn('p:sldIdLst'))
    ids = list(sld_id_lst_now)
    new_id = ids[-1]
    sld_id_lst_now.remove(new_id)
    sld_id_lst_now.insert(insert_pos, new_id)

    print(f'  Inserted timeline slide {src_idx+1} at position {insert_pos+1}')
    return new_slide

# Insert timeline slide 1 at position 3 (after Product Demos, 0-based index 3)
copy_slide_at(timeline, 0, drdo, 3, scale=SCALE)
# Insert timeline slide 2 at position 4 (right after the first timeline slide)
copy_slide_at(timeline, 1, drdo, 4, scale=SCALE)

# ── STEP 3: Save ─────────────────────────────────────────────────────────────
drdo.save(OUT_PATH)
print(f'\nSaved: {OUT_PATH}')
print(f'Total slides: {len(drdo.slides)}')

# Verify order
print('\nFinal slide order:')
for i, slide in enumerate(drdo.slides, 1):
    from lxml import etree
    import re
    texts = re.findall(r'<a:t[^>]*>([^<]+)</a:t>', etree.tostring(slide._element, encoding='unicode'))
    preview = ' | '.join(t.strip() for t in texts if t.strip())[:90]
    print(f'  {i:2d}: {preview}')
